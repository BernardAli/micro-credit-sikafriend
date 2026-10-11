import calendar
import csv
import json
from collections import defaultdict
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO, StringIO

from django.db.models import Sum
from django.http import HttpResponse, JsonResponse
from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from apps.company.models import CompanyProfile

from .models import Borrower, Loan, Payment, Portfolio


QUANT = Decimal("0.01")
RATE_QUANT = Decimal("0.01")

MONTH_CHOICES = [(idx, calendar.month_name[idx]) for idx in range(1, 13)]
QUARTER_CHOICES = [(1, "Q1"), (2, "Q2"), (3, "Q3"), (4, "Q4")]
HALF_YEAR_CHOICES = [(1, "First half"), (2, "Second half")]
PERIOD_CHOICES = [
    ("monthly", "Monthly"),
    ("quarterly", "Quarterly"),
    ("half_yearly", "Half yearly"),
    ("annual", "Annual"),
    ("custom", "Custom range"),
]
EXPORT_FORMATS = [
    ("csv", "CSV"),
    ("excel", "Excel"),
    ("json", "JSON"),
    ("pdf", "PDF"),
]

REPORT_CATALOG = [
    {"key": "portfolio_ownership", "group": "Portfolio", "title": "SikaFriend Portfolio Ownership", "description": "Current portfolio ownership, assigned officers, permanent Customer IDs, and outstanding loans."},
    {"key": "executive_summary", "group": "Management", "title": "Executive Summary", "description": "High-level performance, risk, outreach, and cash movement for the selected window."},
    {"key": "loan_portfolio", "group": "Portfolio", "title": "Loan Portfolio", "description": "Loan book position, repayment status, sector mix, product mix, and outstanding exposure."},
    {"key": "loan_aging", "group": "Portfolio", "title": "Portfolio Aging & PAR", "description": "Current, overdue, and portfolio-at-risk exposure by aging bucket."},
    {"key": "disbursements", "group": "Operations", "title": "Disbursements", "description": "Loans issued in the selected period by borrower, agent, sector, location, and category."},
    {"key": "collections", "group": "Operations", "title": "Collections", "description": "Repayment collections by date, method, agent, borrower, and loan."},
    {"key": "borrower_outreach", "group": "Outreach", "title": "Borrower Outreach", "description": "New clients, active borrowers, gender profile, location spread, and retention indicators."},
    {"key": "agent_performance", "group": "Operations", "title": "Agent Performance", "description": "Field staff disbursement, collection, borrower, and risk contribution."},
    {"key": "sector_breakdown", "group": "Portfolio", "title": "Sector Breakdown", "description": "Portfolio concentration by economic activity and loan purpose sector."},
    {"key": "collateral_register", "group": "Controls", "title": "Collateral Register", "description": "Collateral-backed loans, estimated values, locations, consent, and documentation status."},
    {"key": "mcag_cover", "group": "MCAG / Regulatory", "title": "MCAG Cover Sheet", "description": "Institution and reporting-period details aligned with the supplied MCAG workbook."},
    {"key": "mcag_profit_loss", "group": "MCAG / Regulatory", "title": "MCAG Profit & Loss", "description": "Financial income, operating costs, impairment, and surplus or deficit."},
    {"key": "mcag_balance_sheet", "group": "MCAG / Regulatory", "title": "MCAG Balance Sheet", "description": "Assets, liabilities, equity, loan portfolio, cash, and balancing checks."},
    {"key": "mcag_portfolio", "group": "MCAG / Regulatory", "title": "MCAG Portfolio Report", "description": "Portfolio activity, quality, arrears, provisions, and loan concentration."},
    {"key": "mcag_outreach", "group": "MCAG / Regulatory", "title": "MCAG Outreach & Social Performance", "description": "Borrower outreach, gender, location, client status, and inclusion measures."},
    {"key": "mcag_regional", "group": "MCAG / Regulatory", "title": "MCAG Regional Breakdown", "description": "Regional and location-based borrower, disbursement, collection, and outstanding exposure."},
    {"key": "mcag_ratios", "group": "MCAG / Regulatory", "title": "MCAG Ratios", "description": "Profitability, sustainability, portfolio quality, productivity, and efficiency ratios."},
    {"key": "mcag_workbook", "group": "MCAG / Regulatory", "title": "MCAG Full Workbook", "description": "A multi-sheet Excel workbook covering cover, P&L, balance sheet, portfolio, outreach, regional, and ratios."},
    {"key": "accounting_income_statement", "group": "Accounting", "title": "Income Statement", "description": "Ledger-based income, expenses, and net surplus for the selected period."},
    {"key": "accounting_balance_sheet", "group": "Accounting", "title": "Balance Sheet", "description": "Ledger-based assets, liabilities, equity, and net position as of the report end date."},
    {"key": "accounting_trial_balance", "group": "Accounting", "title": "Trial Balance", "description": "Debit and credit movement by account for the selected reporting window."},
    {"key": "accounting_cashbook", "group": "Accounting", "title": "Cashbook Movement", "description": "Cash and bank debit, credit, movement, and closing balance."},
]


def report_catalog_groups():
    groups = []
    seen = {}
    for report in REPORT_CATALOG:
        group = report["group"]
        if group not in seen:
            seen[group] = {"name": group, "reports": []}
            groups.append(seen[group])
        seen[group]["reports"].append(report)
    return groups


def get_report_definition(key):
    for report in REPORT_CATALOG:
        if report["key"] == key:
            return report
    return REPORT_CATALOG[0]


def available_years(today=None):
    today = today or timezone.localdate()
    return list(range(today.year - 5, today.year + 2))


def _to_int(value, default, minimum=None, maximum=None):
    try:
        number = int(value)
    except (TypeError, ValueError):
        number = default
    if minimum is not None:
        number = max(minimum, number)
    if maximum is not None:
        number = min(maximum, number)
    return number


def _parse_iso_date(value, default):
    try:
        return date.fromisoformat(value) if value else default
    except ValueError:
        return default


def resolve_report_period(params, today=None):
    today = today or timezone.localdate()
    period_type = (params.get("period_type") or "quarterly").strip()
    if period_type not in dict(PERIOD_CHOICES):
        period_type = "quarterly"

    year = _to_int(params.get("year"), today.year, 1900, 2100)
    month = _to_int(params.get("month"), today.month, 1, 12)
    quarter = _to_int(params.get("quarter"), ((today.month - 1) // 3) + 1, 1, 4)
    half = _to_int(params.get("half"), 1 if today.month <= 6 else 2, 1, 2)

    if period_type == "monthly":
        last_day = calendar.monthrange(year, month)[1]
        start = date(year, month, 1)
        end = date(year, month, last_day)
        label = f"{calendar.month_name[month]} {year}"
    elif period_type == "quarterly":
        start_month = ((quarter - 1) * 3) + 1
        end_month = start_month + 2
        start = date(year, start_month, 1)
        end = date(year, end_month, calendar.monthrange(year, end_month)[1])
        label = f"Q{quarter} {year}"
    elif period_type == "half_yearly":
        start_month = 1 if half == 1 else 7
        end_month = 6 if half == 1 else 12
        start = date(year, start_month, 1)
        end = date(year, end_month, calendar.monthrange(year, end_month)[1])
        label = f"H{half} {year}"
    elif period_type == "annual":
        start = date(year, 1, 1)
        end = date(year, 12, 31)
        label = f"FY {year}"
    else:
        default_start = today.replace(day=1)
        default_end = today
        start = _parse_iso_date((params.get("start_date") or "").strip(), default_start)
        end = _parse_iso_date((params.get("end_date") or "").strip(), default_end)
        if start > end:
            start, end = end, start
        label = f"Custom: {start:%d %b %Y} to {end:%d %b %Y}"

    return {
        "type": period_type,
        "year": year,
        "month": month,
        "quarter": quarter,
        "half": half,
        "start": start,
        "end": end,
        "label": label,
    }


def _money(value):
    return Decimal(value or Decimal("0.00")).quantize(QUANT)


def _rate(numerator, denominator):
    numerator = _money(numerator)
    denominator = _money(denominator)
    if denominator == 0:
        return Decimal("0.00")
    return ((numerator / denominator) * Decimal("100")).quantize(RATE_QUANT, rounding=ROUND_HALF_UP)


def _fmt(value):
    if isinstance(value, Decimal):
        return f"{value:,.2f}"
    if isinstance(value, date):
        return value.strftime("%d %b %Y")
    if value is None:
        return ""
    return str(value)


def _metric(label, value, suffix="", help_text=""):
    return {"label": label, "value": f"{_fmt(value)}{suffix}", "help": help_text}


def _section(title, columns, rows, description=""):
    return {"title": title, "description": description, "columns": columns, "rows": [[_fmt(value) for value in row] for row in rows]}


def _empty_section(title, message="No data found for this report period."):
    return _section(title, ["Message"], [[message]])


def _company_profile():
    try:
        return CompanyProfile.get_solo()
    except Exception:
        return None


def _company_value(company, *names, default=""):
    for name in names:
        value = getattr(company, name, None) if company else None
        if value:
            return value
    return default


def _valid_payment(payment):
    return not payment.is_void and not payment.is_pending


def _loan_expected_due_count(loan, as_of):
    if not loan.start_date or as_of < loan.start_date:
        return 0
    if loan.repayment_frequency == "DAILY":
        elapsed = max(0, (as_of - loan.start_date).days)
        return min(loan.expected_payments_count, elapsed)
    elapsed_weeks = max(0, (as_of - loan.start_date).days // 7)
    return min(loan.expected_payments_count, elapsed_weeks)


def _loan_snapshot(loan, period_end):
    payments = [payment for payment in loan.payments.all() if _valid_payment(payment) and payment.payment_date <= period_end]
    paid = _money(sum((payment.amount for payment in payments), Decimal("0.00")))
    principal_paid = _money(sum((payment.principal_paid for payment in payments), Decimal("0.00")))
    roic_paid = _money(sum((payment.roic_paid for payment in payments), Decimal("0.00")))
    expected_total = _money(loan.total_expected_amount)
    principal = _money(loan.principal_amount)
    roic = _money(loan.roic_amount)
    balance = max(expected_total - paid, Decimal("0.00")).quantize(QUANT)
    principal_outstanding = max(principal - principal_paid, Decimal("0.00")).quantize(QUANT)
    roic_outstanding = max(roic - roic_paid, Decimal("0.00")).quantize(QUANT)
    expected_due = _loan_expected_due_count(loan, period_end)
    actual_count = len(payments)
    overdue_units = max(0, expected_due - actual_count)
    return {
        "loan": loan,
        "expected_total": expected_total,
        "principal": principal,
        "roic": roic,
        "paid": paid,
        "principal_paid": principal_paid,
        "roic_paid": roic_paid,
        "balance": balance,
        "principal_outstanding": principal_outstanding,
        "roic_outstanding": roic_outstanding,
        "expected_due": expected_due,
        "actual_count": actual_count,
        "overdue_units": overdue_units,
    }


def _base_snapshot(period):
    start = period["start"]
    end = period["end"]
    loans_qs = Loan.objects.select_related("borrower", "agent_name", "portfolio").prefetch_related("payments")
    loans_as_of = list(loans_qs.filter(start_date__lte=end))
    period_loans = [loan for loan in loans_as_of if loan.start_date and start <= loan.start_date <= end]
    snapshots = [_loan_snapshot(loan, end) for loan in loans_as_of]
    period_snapshots = [item for item in snapshots if item["loan"] in period_loans]
    active_snapshots = [item for item in snapshots if item["balance"] > 0]
    payments = list(
        Payment.objects.select_related("loan", "loan__borrower", "loan__agent_name", "collected_by", "loan__portfolio")
        .filter(is_void=False, is_pending=False, payment_date__gte=start, payment_date__lte=end)
        .order_by("payment_date", "id")
    )
    borrowers_as_of = Borrower.objects.filter(created_at__date__lte=end)
    new_borrowers = borrowers_as_of.filter(created_at__date__gte=start, created_at__date__lte=end)
    return {
        "start": start,
        "end": end,
        "loans_as_of": loans_as_of,
        "period_loans": period_loans,
        "snapshots": snapshots,
        "period_snapshots": period_snapshots,
        "active_snapshots": active_snapshots,
        "payments": payments,
        "borrowers_as_of": borrowers_as_of,
        "new_borrowers": new_borrowers,
    }


def _snapshot_totals(data):
    snapshots = data["snapshots"]
    period_snapshots = data["period_snapshots"]
    active_snapshots = data["active_snapshots"]
    payments = data["payments"]
    total_disbursed = _money(sum((item["principal"] for item in period_snapshots), Decimal("0.00")))
    total_collected = _money(sum((payment.amount for payment in payments), Decimal("0.00")))
    principal_collected = _money(sum((payment.principal_paid for payment in payments), Decimal("0.00")))
    roic_collected = _money(sum((payment.roic_paid for payment in payments), Decimal("0.00")))
    gross_portfolio = _money(sum((item["principal_outstanding"] for item in active_snapshots), Decimal("0.00")))
    total_outstanding = _money(sum((item["balance"] for item in active_snapshots), Decimal("0.00")))
    overdue_outstanding = _money(sum((item["balance"] for item in active_snapshots if item["overdue_units"] > 0), Decimal("0.00")))
    par_6 = _money(sum((item["balance"] for item in active_snapshots if item["overdue_units"] > 6), Decimal("0.00")))
    return {
        "total_disbursed": total_disbursed,
        "total_collected": total_collected,
        "principal_collected": principal_collected,
        "roic_collected": roic_collected,
        "gross_portfolio": gross_portfolio,
        "total_outstanding": total_outstanding,
        "overdue_outstanding": overdue_outstanding,
        "par_6": par_6,
        "collection_rate": _rate(total_collected, sum((item["expected_total"] for item in period_snapshots), Decimal("0.00"))),
        "repayment_rate": _rate(total_collected, total_disbursed),
        "par_6_rate": _rate(par_6, total_outstanding),
    }


def _status_label(loan, snapshot):
    if snapshot["balance"] <= 0:
        return "Closed"
    if snapshot["overdue_units"] > 6:
        return "Default risk"
    if snapshot["overdue_units"] > 0:
        return "Overdue"
    return "Active"


def _sector_label(loan):
    if loan.sector == "OTHER" and loan.sector_other:
        return loan.sector_other
    return loan.get_sector_display() if loan.sector else "Unspecified"


def _location_label(loan):
    return (loan.borrower.location or "Unassigned").strip() if loan.borrower_id else "Unassigned"


def _agent_label(loan):
    return loan.agent_name.name if loan.agent_name_id else "Unassigned"


def _accounting_available_report(start, end):
    try:
        from apps.accounting.models import JournalLine, LedgerAccount
        period_lines = JournalLine.objects.select_related("account", "entry").filter(
            entry__status="POSTED",
            entry__entry_date__gte=start,
            entry__entry_date__lte=end,
        )
        as_of_lines = JournalLine.objects.select_related("account", "entry").filter(
            entry__status="POSTED",
            entry__entry_date__lte=end,
        )
        accounts = list(LedgerAccount.objects.filter(is_active=True).order_by("code", "name"))
    except Exception as exc:
        return {"available": False, "error": str(exc), "period_rows": [], "as_of_rows": []}

    def rows_for(lines):
        debit_credit = defaultdict(lambda: {"debit": Decimal("0.00"), "credit": Decimal("0.00")})
        for line in lines:
            row = debit_credit[line.account_id]
            row["debit"] += _money(line.debit)
            row["credit"] += _money(line.credit)
        rows = []
        for account in accounts:
            values = debit_credit.get(account.pk, {"debit": Decimal("0.00"), "credit": Decimal("0.00")})
            debit = _money(values["debit"])
            credit = _money(values["credit"])
            if account.account_type in {"ASSET", "EXPENSE"}:
                balance = _money(debit - credit)
            else:
                balance = _money(credit - debit)
            if debit == 0 and credit == 0 and not account.is_system:
                continue
            rows.append({"account": account, "debit": debit, "credit": credit, "balance": balance})
        return rows

    return {"available": True, "error": "", "period_rows": rows_for(period_lines), "as_of_rows": rows_for(as_of_lines)}


def _accounting_totals(start, end):
    data = _accounting_available_report(start, end)
    totals = defaultdict(lambda: Decimal("0.00"))
    if not data["available"]:
        return data, totals
    for row in data["period_rows"]:
        account_type = row["account"].account_type
        if account_type in {"INCOME", "EXPENSE"}:
            totals[account_type] += row["balance"]
    for row in data["as_of_rows"]:
        account_type = row["account"].account_type
        if account_type in {"ASSET", "LIABILITY", "EQUITY"}:
            totals[account_type] += row["balance"]
    for key in list(totals):
        totals[key] = _money(totals[key])
    return data, totals


def _report_shell(key, period, metrics, sections):
    definition = get_report_definition(key)
    return {
        "key": key,
        "title": definition["title"],
        "description": definition["description"],
        "period": period,
        "metrics": metrics,
        "sections": sections,
    }


def _executive_summary(key, period, data):
    totals = _snapshot_totals(data)
    accounting, accounting_totals = _accounting_totals(data["start"], data["end"])
    income = accounting_totals["INCOME"] or totals["roic_collected"]
    expense = accounting_totals["EXPENSE"]
    net_surplus = _money(income - expense)
    metrics = [
        _metric("Loans disbursed", len(data["period_loans"])),
        _metric("Disbursed amount", totals["total_disbursed"]),
        _metric("Collections", totals["total_collected"]),
        _metric("Portfolio outstanding", totals["total_outstanding"]),
        _metric("PAR > 6", totals["par_6"], " GHS", "Outstanding amount more than six expected repayments behind."),
        _metric("Net surplus", net_surplus),
    ]
    sections = [
        _section("Performance Summary", ["Measure", "Value"], [
            ["New borrowers", data["new_borrowers"].count()],
            ["Borrowers as of period end", data["borrowers_as_of"].count()],
            ["Active outstanding loans", len(data["active_snapshots"])],
            ["Gross principal portfolio", totals["gross_portfolio"]],
            ["ROIC collected", totals["roic_collected"]],
            ["Principal collected", totals["principal_collected"]],
            ["Collection rate", f"{totals['collection_rate']}%"],
            ["PAR > 6 rate", f"{totals['par_6_rate']}%"],
        ]),
        _section("Management Attention", ["Area", "Indicator"], [
            ["Portfolio risk", f"{len([x for x in data['active_snapshots'] if x['overdue_units'] > 0])} outstanding loans have overdue repayments."],
            ["Liquidity movement", f"GHS {totals['total_collected']} collected against GHS {totals['total_disbursed']} disbursed."],
            ["Accounting basis", "Ledger accounts used where available; loan ROIC collections used as income fallback." if not accounting["available"] else "Posted ledger entries used for financial lines."],
        ]),
    ]
    return _report_shell(key, period, metrics, sections)


def _loan_portfolio(key, period, data):
    totals = _snapshot_totals(data)
    status_rows = defaultdict(lambda: {"count": 0, "principal": Decimal("0.00"), "balance": Decimal("0.00")})
    category_rows = defaultdict(lambda: {"count": 0, "principal": Decimal("0.00"), "balance": Decimal("0.00")})
    loan_rows = []
    for item in data["snapshots"]:
        loan = item["loan"]
        status = _status_label(loan, item)
        status_rows[status]["count"] += 1
        status_rows[status]["principal"] += item["principal"]
        status_rows[status]["balance"] += item["balance"]
        category = loan.get_loan_category_display() if loan.loan_category else "Unspecified"
        category_rows[category]["count"] += 1
        category_rows[category]["principal"] += item["principal"]
        category_rows[category]["balance"] += item["balance"]
        if item["balance"] > 0:
            loan_rows.append([
                loan.id,
                f"{loan.borrower.full_name} ({loan.borrower.borrower_id})",
                _agent_label(loan),
                category,
                _sector_label(loan),
                loan.start_date,
                item["principal"],
                item["paid"],
                item["balance"],
                item["overdue_units"],
                status,
            ])
    sections = [
        _section("Status Summary", ["Status", "Loans", "Principal", "Outstanding"], [
            [name, row["count"], _money(row["principal"]), _money(row["balance"])]
            for name, row in sorted(status_rows.items())
        ]),
        _section("Category Summary", ["Category", "Loans", "Principal", "Outstanding"], [
            [name, row["count"], _money(row["principal"]), _money(row["balance"])]
            for name, row in sorted(category_rows.items())
        ]),
        _section("Outstanding Loan Book", ["Loan", "Borrower", "Agent", "Category", "Sector", "Start Date", "Principal", "Paid", "Outstanding", "Overdue Units", "Status"], loan_rows[:500]) if loan_rows else _empty_section("Outstanding Loan Book"),
    ]
    metrics = [
        _metric("Loans as of end", len(data["snapshots"])),
        _metric("Active balances", len(data["active_snapshots"])),
        _metric("Gross portfolio", totals["gross_portfolio"]),
        _metric("Total outstanding", totals["total_outstanding"]),
    ]
    return _report_shell(key, period, metrics, sections)


def _loan_aging(key, period, data):
    buckets = [
        ("current", "Current", lambda value: value == 0),
        ("1_4", "1-4 overdue units", lambda value: 1 <= value <= 4),
        ("5_8", "5-8 overdue units", lambda value: 5 <= value <= 8),
        ("9_12", "9-12 overdue units", lambda value: 9 <= value <= 12),
        ("12_plus", "Over 12 overdue units", lambda value: value > 12),
    ]
    bucket_rows = []
    for _, label, predicate in buckets:
        matches = [item for item in data["active_snapshots"] if predicate(item["overdue_units"])]
        balance = _money(sum((item["balance"] for item in matches), Decimal("0.00")))
        principal = _money(sum((item["principal_outstanding"] for item in matches), Decimal("0.00")))
        bucket_rows.append([label, len(matches), principal, balance])
    watch_rows = sorted(
        [
            [
                item["loan"].id,
                item["loan"].borrower.full_name,
                _agent_label(item["loan"]),
                _location_label(item["loan"]),
                item["loan"].start_date,
                item["balance"],
                item["overdue_units"],
                "PAR > 6" if item["overdue_units"] > 6 else "Watch",
            ]
            for item in data["active_snapshots"]
            if item["overdue_units"] > 0
        ],
        key=lambda row: (-int(row[6]), Decimal(row[5].replace(",", "")) if isinstance(row[5], str) else -row[5]),
    )
    totals = _snapshot_totals(data)
    metrics = [
        _metric("Overdue outstanding", totals["overdue_outstanding"]),
        _metric("PAR > 6", totals["par_6"]),
        _metric("PAR > 6 rate", totals["par_6_rate"], "%"),
        _metric("Watch loans", len(watch_rows)),
    ]
    sections = [
        _section("Aging Buckets", ["Bucket", "Loans", "Principal Outstanding", "Total Outstanding"], bucket_rows),
        _section("Overdue Watch List", ["Loan", "Borrower", "Agent", "Location", "Start Date", "Outstanding", "Overdue Units", "Risk"], watch_rows[:500]) if watch_rows else _empty_section("Overdue Watch List"),
    ]
    return _report_shell(key, period, metrics, sections)


def _disbursements(key, period, data):
    rows = [
        [
            loan.id,
            loan.start_date,
            f"{loan.borrower.full_name} ({loan.borrower.borrower_id})",
            _agent_label(loan),
            _location_label(loan),
            _sector_label(loan),
            loan.get_loan_category_display() if loan.loan_category else "",
            loan.repayment_frequency,
            loan.amount,
            loan.weekly_payment,
            loan.expected_payments_count,
        ]
        for loan in sorted(data["period_loans"], key=lambda item: (item.start_date, item.id))
    ]
    totals = _snapshot_totals(data)
    by_agent = defaultdict(lambda: {"count": 0, "amount": Decimal("0.00")})
    for loan in data["period_loans"]:
        row = by_agent[_agent_label(loan)]
        row["count"] += 1
        row["amount"] += _money(loan.amount)
    sections = [
        _section("Disbursement Summary by Agent", ["Agent", "Loans", "Amount"], [[name, row["count"], _money(row["amount"])] for name, row in sorted(by_agent.items())]),
        _section("Disbursement Register", ["Loan", "Start Date", "Borrower", "Agent", "Location", "Sector", "Category", "Frequency", "Principal", "Installment", "Expected Payments"], rows[:1000]) if rows else _empty_section("Disbursement Register"),
    ]
    return _report_shell(key, period, [_metric("Loans issued", len(rows)), _metric("Amount issued", totals["total_disbursed"])], sections)


def _collections(key, period, data):
    payments = data["payments"]
    by_method = defaultdict(lambda: {"count": 0, "amount": Decimal("0.00"), "principal": Decimal("0.00"), "roic": Decimal("0.00")})
    by_agent = defaultdict(lambda: {"count": 0, "amount": Decimal("0.00")})
    rows = []
    for payment in payments:
        method = payment.get_payment_method_display()
        by_method[method]["count"] += 1
        by_method[method]["amount"] += _money(payment.amount)
        by_method[method]["principal"] += _money(payment.principal_paid)
        by_method[method]["roic"] += _money(payment.roic_paid)
        by_agent[_collection_officer(payment)]["count"] += 1
        by_agent[_collection_officer(payment)]["amount"] += _money(payment.amount)
        rows.append([
            payment.payment_date,
            payment.id,
            payment.loan_id,
            f"{payment.loan.borrower.full_name} ({payment.loan.borrower.borrower_id})",
            _collection_officer(payment),
            method,
            payment.reference,
            payment.amount,
            payment.principal_paid,
            payment.roic_paid,
        ])
    total = _money(sum((payment.amount for payment in payments), Decimal("0.00")))
    sections = [
        _section("Collections by Method", ["Method", "Payments", "Amount", "Principal", "ROIC"], [[name, row["count"], _money(row["amount"]), _money(row["principal"]), _money(row["roic"])] for name, row in sorted(by_method.items())]),
        _section("Collections by Agent", ["Agent", "Payments", "Amount"], [[name, row["count"], _money(row["amount"])] for name, row in sorted(by_agent.items())]),
        _section("Collection Register", ["Date", "Payment", "Loan", "Borrower", "Agent", "Method", "Reference", "Amount", "Principal", "ROIC"], rows[:1000]) if rows else _empty_section("Collection Register"),
    ]
    return _report_shell(key, period, [_metric("Payments", len(payments)), _metric("Collections", total)], sections)


def _collection_officer(payment):
    return payment.collected_by.name if payment.collected_by_id else "Staff / unrecorded collector"


def _portfolio_ownership(key, period, data):
    totals = defaultdict(lambda: {"loans": 0, "outstanding": Decimal("0.00"), "collections": Decimal("0.00")})
    register = []
    for snapshot in data["snapshots"]:
        loan = snapshot["loan"]
        row = totals[loan.portfolio_id]
        row["loans"] += 1
        row["outstanding"] += snapshot["balance"]
        register.append([loan.borrower.borrower_id, loan.borrower.legacy_borrower_id, f"{loan.borrower.full_name} ({loan.borrower.borrower_id})",
                         str(loan.portfolio or "Needs assignment"), loan.pk, snapshot["balance"]])
    for payment in data["payments"]:
        totals[payment.loan.portfolio_id]["collections"] += payment.amount
    summary = []
    for portfolio in Portfolio.objects.prefetch_related("assignments__agent"):
        row = totals[portfolio.pk]
        officers = ", ".join(assignment.agent.name for assignment in portfolio.assignments.all() if assignment.ended_at is None)
        summary.append([portfolio.number, portfolio.name, officers or "Needs assignment", portfolio.borrowers.count(),
                        row["loans"], _money(row["outstanding"]), _money(row["collections"])])
    pending = totals[None]
    if pending["loans"]:
        summary.append(["", "Needs assignment", "", Borrower.objects.filter(portfolio__isnull=True).count(), pending["loans"], _money(pending["outstanding"]), _money(pending["collections"])])
    return _report_shell(key, period, [_metric("Named portfolios", Portfolio.objects.count())], [
        _section("Current Portfolio Ownership", ["Portfolio No.", "Name", "Current Officers", "Customers", "Loans", "Outstanding", "Period Collections"], summary),
        _section("Customer Loan Register", ["Customer ID", "Legacy ID", "Customer", "Portfolio", "Loan", "Outstanding"], register),
    ])


def _borrower_outreach(key, period, data):
    borrowers = data["borrowers_as_of"]
    gender_rows = []
    for code, label in Borrower.GENDER_CHOICES:
        gender_rows.append([label, borrowers.filter(gender=code).count(), data["new_borrowers"].filter(gender=code).count()])
    gender_rows.append(["Unspecified", borrowers.filter(gender="").count(), data["new_borrowers"].filter(gender="").count()])
    location_rows = []
    for location in borrowers.values_list("location", flat=True).distinct().order_by("location"):
        name = location or "Unassigned"
        active_count = borrowers.filter(location=location).count()
        new_count = data["new_borrowers"].filter(location=location).count()
        outstanding = _money(sum((item["balance"] for item in data["active_snapshots"] if _location_label(item["loan"]) == name), Decimal("0.00")))
        location_rows.append([name, active_count, new_count, outstanding])
    metrics = [
        _metric("Borrowers as of end", borrowers.count()),
        _metric("New borrowers", data["new_borrowers"].count()),
        _metric("Borrowers with active balances", len({item["loan"].borrower_id for item in data["active_snapshots"]})),
    ]
    sections = [
        _section("Gender Outreach", ["Gender", "Total Borrowers", "New Borrowers"], gender_rows),
        _section("Location Outreach", ["Location", "Total Borrowers", "New Borrowers", "Outstanding"], location_rows[:500]) if location_rows else _empty_section("Location Outreach"),
    ]
    return _report_shell(key, period, metrics, sections)


def _agent_performance(key, period, data):
    agent_rows = defaultdict(lambda: {
        "period_loans": 0,
        "period_disbursed": Decimal("0.00"),
        "collections": Decimal("0.00"),
        "borrowers": set(),
        "outstanding": Decimal("0.00"),
        "overdue": Decimal("0.00"),
    })
    for loan in data["period_loans"]:
        row = agent_rows[_agent_label(loan)]
        row["period_loans"] += 1
        row["period_disbursed"] += _money(loan.amount)
        row["borrowers"].add(loan.borrower_id)
    for payment in data["payments"]:
        agent_rows[_collection_officer(payment)]["collections"] += _money(payment.amount)
    for item in data["active_snapshots"]:
        row = agent_rows[_agent_label(item["loan"])]
        row["borrowers"].add(item["loan"].borrower_id)
        row["outstanding"] += item["balance"]
        if item["overdue_units"] > 0:
            row["overdue"] += item["balance"]
    rows = [
        [
            name,
            row["period_loans"],
            len(row["borrowers"]),
            _money(row["period_disbursed"]),
            _money(row["collections"]),
            _money(row["outstanding"]),
            _money(row["overdue"]),
            f"{_rate(row['collections'], row['period_disbursed'])}%",
        ]
        for name, row in sorted(agent_rows.items())
    ]
    return _report_shell(key, period, [_metric("Agents reported", len(rows))], [_section("Agent Performance", ["Agent", "Loans Issued", "Borrowers", "Disbursed", "Collected", "Outstanding", "Overdue", "Collection/Disbursement"], rows) if rows else _empty_section("Agent Performance")])


def _sector_breakdown(key, period, data):
    sector_rows = defaultdict(lambda: {"loans": 0, "borrowers": set(), "disbursed": Decimal("0.00"), "outstanding": Decimal("0.00"), "overdue": Decimal("0.00")})
    for item in data["snapshots"]:
        loan = item["loan"]
        row = sector_rows[_sector_label(loan)]
        row["loans"] += 1
        row["borrowers"].add(loan.borrower_id)
        row["outstanding"] += item["balance"]
        if item["overdue_units"] > 0:
            row["overdue"] += item["balance"]
    for loan in data["period_loans"]:
        sector_rows[_sector_label(loan)]["disbursed"] += _money(loan.amount)
    rows = [
        [name, row["loans"], len(row["borrowers"]), _money(row["disbursed"]), _money(row["outstanding"]), _money(row["overdue"]), f"{_rate(row['overdue'], row['outstanding'])}%"]
        for name, row in sorted(sector_rows.items())
    ]
    return _report_shell(key, period, [_metric("Sectors", len(rows))], [_section("Sector Concentration", ["Sector", "Loans", "Borrowers", "Disbursed", "Outstanding", "Overdue", "Overdue Rate"], rows) if rows else _empty_section("Sector Concentration")])


def _collateral_register(key, period, data):
    rows = []
    collateral_value = Decimal("0.00")
    secured_balance = Decimal("0.00")
    for item in data["active_snapshots"]:
        loan = item["loan"]
        if not loan.collateral_type and not loan.collateral_details and not loan.collateral_estimated_value:
            continue
        value = _money(loan.collateral_estimated_value)
        collateral_value += value
        secured_balance += item["balance"]
        rows.append([
            loan.id,
            f"{loan.borrower.full_name} ({loan.borrower.borrower_id})",
            _agent_label(loan),
            loan.collateral_type or "Unspecified",
            loan.collateral_details,
            loan.collateral_location,
            value,
            item["balance"],
            "Yes" if loan.collateral_borrower_consent_accepted else "No",
            loan.collateral_borrower_signed_date,
            loan.collateral_guarantor_signature_name,
        ])
    metrics = [
        _metric("Secured active loans", len(rows)),
        _metric("Collateral value", _money(collateral_value)),
        _metric("Secured outstanding", _money(secured_balance)),
        _metric("Coverage", _rate(collateral_value, secured_balance), "%"),
    ]
    return _report_shell(key, period, metrics, [_section("Collateral Register", ["Loan", "Borrower", "Agent", "Type", "Details", "Location", "Estimated Value", "Outstanding", "Consent", "Signed Date", "Guarantor"], rows[:1000]) if rows else _empty_section("Collateral Register")])


def _mcag_cover(key, period, data):
    company = _company_profile()
    quarter = ((period["start"].month - 1) // 3) + 1
    rows = [
        ["Institution / Member Name", _company_value(company, "name", "company_name", default="")],
        ["Member Code", _company_value(company, "member_code", "license_number", "registration_number", default="")],
        ["Return Code", "MCAG"],
        ["Reporting Year", period["start"].year],
        ["Reporting Quarter", f"Q{quarter}"],
        ["Reporting Period", period["label"]],
        ["Reporting Start Date", period["start"]],
        ["Reporting End Date", period["end"]],
        ["Region", _company_value(company, "region", "location", default="")],
        ["Generated On", timezone.localdate()],
    ]
    return _report_shell(key, period, [_metric("Report period", period["label"])], [_section("Cover Sheet", ["Field", "Value"], rows)])


def _mcag_profit_loss(key, period, data):
    totals = _snapshot_totals(data)
    accounting, accounting_totals = _accounting_totals(data["start"], data["end"])
    income = accounting_totals["INCOME"] or totals["roic_collected"]
    expenses = accounting_totals["EXPENSE"]
    impairment = totals["par_6"] * Decimal("0.10")
    surplus = _money(income - expenses - impairment)
    rows = [
        ["Financial income from loan portfolio", income],
        ["Loan ROIC collected", totals["roic_collected"]],
        ["Other operating income", max(income - totals["roic_collected"], Decimal("0.00"))],
        ["Operating expenses", expenses],
        ["Loan loss provision estimate", _money(impairment)],
        ["Surplus / (deficit)", surplus],
    ]
    if not accounting["available"]:
        rows.append(["Accounting note", "Ledger data unavailable; income falls back to loan ROIC collections and expense lines remain zero."])
    metrics = [_metric("Income", income), _metric("Expenses", expenses), _metric("Surplus", surplus)]
    return _report_shell(key, period, metrics, [_section("Profit & Loss Statement", ["Line Item", "Amount"], rows)])


def _mcag_balance_sheet(key, period, data):
    totals = _snapshot_totals(data)
    accounting, accounting_totals = _accounting_totals(data["start"], data["end"])
    assets = accounting_totals["ASSET"] or totals["gross_portfolio"]
    liabilities = accounting_totals["LIABILITY"]
    equity = accounting_totals["EQUITY"]
    net_position = _money(assets - liabilities - equity)
    rows = [
        ["Cash and bank balances", max(assets - totals["gross_portfolio"], Decimal("0.00"))],
        ["Gross loan portfolio principal", totals["gross_portfolio"]],
        ["Loan ROIC outstanding", _money(sum((item["roic_outstanding"] for item in data["active_snapshots"]), Decimal("0.00")))],
        ["Total assets", assets],
        ["Total liabilities", liabilities],
        ["Total equity", equity],
        ["Balance check", net_position],
    ]
    if not accounting["available"]:
        rows.append(["Accounting note", "Ledger data unavailable; assets are estimated from the loan portfolio."])
    metrics = [_metric("Assets", assets), _metric("Liabilities", liabilities), _metric("Equity", equity), _metric("Balance check", net_position)]
    return _report_shell(key, period, metrics, [_section("Balance Sheet", ["Line Item", "Amount"], rows)])


def _mcag_portfolio(key, period, data):
    totals = _snapshot_totals(data)
    rows = [
        ["Number of loans disbursed during period", len(data["period_loans"])],
        ["Value of loans disbursed during period", totals["total_disbursed"]],
        ["Number of active outstanding loans", len(data["active_snapshots"])],
        ["Gross outstanding principal portfolio", totals["gross_portfolio"]],
        ["Total outstanding including ROIC", totals["total_outstanding"]],
        ["Principal collected during period", totals["principal_collected"]],
        ["ROIC collected during period", totals["roic_collected"]],
        ["Outstanding overdue portfolio", totals["overdue_outstanding"]],
        ["PAR greater than six units", totals["par_6"]],
        ["PAR greater than six units rate", f"{totals['par_6_rate']}%"],
        ["Estimated loan loss provision at 10% of PAR > 6", _money(totals["par_6"] * Decimal("0.10"))],
    ]
    metrics = [_metric("Disbursed", totals["total_disbursed"]), _metric("Outstanding", totals["total_outstanding"]), _metric("PAR > 6", totals["par_6"]), _metric("PAR rate", totals["par_6_rate"], "%")]
    return _report_shell(key, period, metrics, [_section("Portfolio Report", ["Indicator", "Value"], rows)])


def _mcag_outreach(key, period, data):
    report = _borrower_outreach(key, period, data)
    report["title"] = get_report_definition(key)["title"]
    report["description"] = get_report_definition(key)["description"]
    return report


def _mcag_regional(key, period, data):
    regional = defaultdict(lambda: {"borrowers": set(), "new_borrowers": 0, "loans": 0, "disbursed": Decimal("0.00"), "collections": Decimal("0.00"), "outstanding": Decimal("0.00"), "overdue": Decimal("0.00")})
    for borrower in data["borrowers_as_of"]:
        regional[borrower.location or "Unassigned"]["borrowers"].add(borrower.pk)
    for borrower in data["new_borrowers"]:
        regional[borrower.location or "Unassigned"]["new_borrowers"] += 1
    for loan in data["period_loans"]:
        row = regional[_location_label(loan)]
        row["loans"] += 1
        row["disbursed"] += _money(loan.amount)
    for payment in data["payments"]:
        regional[_location_label(payment.loan)]["collections"] += _money(payment.amount)
    for item in data["active_snapshots"]:
        row = regional[_location_label(item["loan"])]
        row["outstanding"] += item["balance"]
        if item["overdue_units"] > 0:
            row["overdue"] += item["balance"]
    rows = [
        [name, len(row["borrowers"]), row["new_borrowers"], row["loans"], _money(row["disbursed"]), _money(row["collections"]), _money(row["outstanding"]), _money(row["overdue"])]
        for name, row in sorted(regional.items())
    ]
    return _report_shell(key, period, [_metric("Locations", len(rows))], [_section("Regional Breakdown", ["Location / Region", "Borrowers", "New Borrowers", "Loans Disbursed", "Disbursed", "Collections", "Outstanding", "Overdue"], rows) if rows else _empty_section("Regional Breakdown")])


def _mcag_ratios(key, period, data):
    totals = _snapshot_totals(data)
    accounting, accounting_totals = _accounting_totals(data["start"], data["end"])
    income = accounting_totals["INCOME"] or totals["roic_collected"]
    expenses = accounting_totals["EXPENSE"]
    assets = accounting_totals["ASSET"] or totals["gross_portfolio"]
    equity = accounting_totals["EQUITY"]
    active_borrowers = len({item["loan"].borrower_id for item in data["active_snapshots"]})
    rows = [
        ["Operational self-sufficiency", f"{_rate(income, expenses)}%", "Financial income divided by operating expenses."],
        ["Return on assets", f"{_rate(income - expenses, assets)}%", "Surplus divided by total assets."],
        ["Return on equity", f"{_rate(income - expenses, equity)}%", "Surplus divided by equity."],
        ["Portfolio yield", f"{_rate(totals['roic_collected'], totals['gross_portfolio'])}%", "ROIC collected divided by gross loan portfolio."],
        ["Repayment rate", f"{totals['repayment_rate']}%", "Collections divided by disbursements in the selected period."],
        ["PAR > 6 rate", f"{totals['par_6_rate']}%", "Portfolio at risk over six expected repayments divided by total outstanding."],
        ["Average loan size", _money(totals["total_disbursed"] / Decimal(len(data["period_loans"]))) if data["period_loans"] else Decimal("0.00"), "Disbursements divided by number of loans issued."],
        ["Borrowers per active loan", _money(Decimal(active_borrowers) / Decimal(len(data["active_snapshots"]))) if data["active_snapshots"] else Decimal("0.00"), "Active borrowers divided by active outstanding loans."],
    ]
    if not accounting["available"]:
        rows.append(["Accounting note", "Ledger data unavailable", "Profitability ratios use loan income fallback and may be incomplete."])
    return _report_shell(key, period, [_metric("Ratios generated", len(rows))], [_section("Ratios", ["Ratio", "Value", "Definition"], rows)])


def _accounting_report(key, period, data):
    accounting, totals = _accounting_totals(data["start"], data["end"])
    if not accounting["available"]:
        return _report_shell(key, period, [_metric("Status", "Unavailable")], [_empty_section("Accounting Report", accounting["error"] or "Accounting data is unavailable.")])
    if key == "accounting_income_statement":
        rows = [[row["account"].code, row["account"].name, row["account"].account_type, row["balance"]] for row in accounting["period_rows"] if row["account"].account_type in {"INCOME", "EXPENSE"}]
        rows.append(["", "Net Surplus / (Deficit)", "", _money(totals["INCOME"] - totals["EXPENSE"])])
        metrics = [_metric("Income", totals["INCOME"]), _metric("Expenses", totals["EXPENSE"]), _metric("Net surplus", _money(totals["INCOME"] - totals["EXPENSE"]))]
        title = "Income Statement"
    elif key == "accounting_balance_sheet":
        rows = [[row["account"].code, row["account"].name, row["account"].account_type, row["balance"]] for row in accounting["as_of_rows"] if row["account"].account_type in {"ASSET", "LIABILITY", "EQUITY"}]
        rows.append(["", "Balance Check", "", _money(totals["ASSET"] - totals["LIABILITY"] - totals["EQUITY"])])
        metrics = [_metric("Assets", totals["ASSET"]), _metric("Liabilities", totals["LIABILITY"]), _metric("Equity", totals["EQUITY"])]
        title = "Balance Sheet"
    elif key == "accounting_trial_balance":
        rows = [[row["account"].code, row["account"].name, row["debit"], row["credit"], row["balance"]] for row in accounting["period_rows"]]
        metrics = [_metric("Accounts", len(rows)), _metric("Debits", sum((row["debit"] for row in accounting["period_rows"]), Decimal("0.00"))), _metric("Credits", sum((row["credit"] for row in accounting["period_rows"]), Decimal("0.00")))]
        return _report_shell(key, period, metrics, [_section("Trial Balance", ["Code", "Account", "Debit", "Credit", "Balance"], rows)])
    else:
        rows = [[row["account"].code, row["account"].name, row["debit"], row["credit"], row["balance"]] for row in accounting["period_rows"] if getattr(row["account"], "is_cash_account", False) or getattr(row["account"], "is_bank_account", False)]
        metrics = [_metric("Cash accounts", len(rows))]
        return _report_shell(key, period, metrics, [_section("Cashbook Movement", ["Code", "Account", "Debit", "Credit", "Movement"], rows) if rows else _empty_section("Cashbook Movement")])
    return _report_shell(key, period, metrics, [_section(title, ["Code", "Account", "Type", "Balance"], rows) if rows else _empty_section(title)])


REPORT_BUILDERS = {
    "portfolio_ownership": _portfolio_ownership,
    "executive_summary": _executive_summary,
    "loan_portfolio": _loan_portfolio,
    "loan_aging": _loan_aging,
    "disbursements": _disbursements,
    "collections": _collections,
    "borrower_outreach": _borrower_outreach,
    "agent_performance": _agent_performance,
    "sector_breakdown": _sector_breakdown,
    "collateral_register": _collateral_register,
    "mcag_cover": _mcag_cover,
    "mcag_profit_loss": _mcag_profit_loss,
    "mcag_balance_sheet": _mcag_balance_sheet,
    "mcag_portfolio": _mcag_portfolio,
    "mcag_outreach": _mcag_outreach,
    "mcag_regional": _mcag_regional,
    "mcag_ratios": _mcag_ratios,
    "accounting_income_statement": _accounting_report,
    "accounting_balance_sheet": _accounting_report,
    "accounting_trial_balance": _accounting_report,
    "accounting_cashbook": _accounting_report,
}

MCAG_WORKBOOK_KEYS = [
    "mcag_cover",
    "mcag_profit_loss",
    "mcag_balance_sheet",
    "mcag_portfolio",
    "mcag_outreach",
    "mcag_regional",
    "mcag_ratios",
]


def build_report(key, period):
    key = get_report_definition(key)["key"]
    data = _base_snapshot(period)
    if key == "mcag_workbook":
        reports = [REPORT_BUILDERS[report_key](report_key, period, data) for report_key in MCAG_WORKBOOK_KEYS]
        metrics = [
            _metric("Sheets", len(reports)),
            _metric("Period", period["label"]),
        ]
        sections = []
        for report in reports:
            sections.extend(report["sections"])
        workbook_report = _report_shell(key, period, metrics, sections)
        workbook_report["child_reports"] = reports
        return workbook_report
    builder = REPORT_BUILDERS.get(key, _executive_summary)
    return builder(key, period, data)


def _filename(report, extension):
    return f"{slugify(report['title']) or 'report'}-{report['period']['start']:%Y%m%d}-{report['period']['end']:%Y%m%d}.{extension}"


def export_report_csv(report):
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow([report["title"], report["period"]["label"]])
    writer.writerow(["Start", report["period"]["start"], "End", report["period"]["end"]])
    writer.writerow([])
    writer.writerow(["Metric", "Value"])
    for metric in report["metrics"]:
        writer.writerow([metric["label"], metric["value"]])
    for section in report["sections"]:
        writer.writerow([])
        writer.writerow([section["title"]])
        if section.get("description"):
            writer.writerow([section["description"]])
        writer.writerow(section["columns"])
        writer.writerows(section["rows"])
    response = HttpResponse(output.getvalue(), content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{_filename(report, "csv")}"'
    return response


def _append_report_to_sheet(ws, report):
    ws.append([report["title"]])
    ws.append(["Period", report["period"]["label"], "Start", report["period"]["start"], "End", report["period"]["end"]])
    ws.append([])
    ws.append(["Metric", "Value"])
    for metric in report["metrics"]:
        ws.append([metric["label"], metric["value"]])
    for section in report["sections"]:
        ws.append([])
        ws.append([section["title"]])
        ws.append(section["columns"])
        for row in section["rows"]:
            ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
    for column_cells in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 42)


def export_report_excel(report, period=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    if report["key"] == "mcag_workbook":
        wb.remove(ws)
        for child in report.get("child_reports") or [build_report(report_key, period or report["period"]) for report_key in MCAG_WORKBOOK_KEYS]:
            title = child["title"].replace("MCAG ", "")[:31]
            sheet = wb.create_sheet(title=title)
            _append_report_to_sheet(sheet, child)
    else:
        _append_report_to_sheet(ws, report)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    response = HttpResponse(output.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{_filename(report, "xlsx")}"'
    return response


def export_report_json(report):
    response = JsonResponse(json.loads(json.dumps(report, default=str)), json_dumps_params={"indent": 2})
    response["Content-Disposition"] = f'attachment; filename="{_filename(report, "json")}"'
    return response
