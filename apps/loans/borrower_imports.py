from __future__ import annotations

import csv
import io
import json
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.db import transaction
from openpyxl import Workbook, load_workbook

from apps.tenants.enforcement import can_create_borrower

from .models import Borrower

UserModel = get_user_model()

BORROWER_IMPORT_FORMATS = [
    {
        "slug": "csv",
        "label": "CSV Spreadsheet",
        "extensions": ".csv",
        "description": "Simple comma-separated rows for spreadsheet tools and quick data preparation.",
    },
    {
        "slug": "excel",
        "label": "Excel Workbook",
        "extensions": ".xlsx",
        "description": "Structured worksheet import for operations teams working in Microsoft Excel.",
    },
    {
        "slug": "json",
        "label": "JSON Payload",
        "extensions": ".json",
        "description": "List-based structured payload for integrations, scripts, and data pipelines.",
    },
    {
        "slug": "xml",
        "label": "XML Document",
        "extensions": ".xml",
        "description": "Tag-based import format for legacy data exchange and structured archives.",
    },
]

BORROWER_SAMPLE_COLUMNS = [
    "borrower_id",
    "full_name",
    "national_id",
    "phone",
    "email",
    "location",
    "ghana_card_number",
    "group_name",
    "preferred_language",
    "status",
]


@dataclass
class BorrowerImportError:
    row_number: int
    message: str
    row_data: dict


def get_borrower_import_formats():
    return BORROWER_IMPORT_FORMATS


def get_borrower_import_column_help():
    return [
        {"name": "borrower_id", "required": True, "help": "Unique borrower code used in the platform."},
        {"name": "full_name", "required": True, "help": "Borrower full name as it should appear on records."},
        {"name": "national_id", "required": True, "help": "Unique national identity reference."},
        {"name": "phone", "required": True, "help": "Primary mobile or contact number."},
        {"name": "email", "required": False, "help": "Optional recovery email for future portal access."},
        {"name": "location", "required": True, "help": "Town, community, or residential area."},
        {"name": "ghana_card_number", "required": True, "help": "Unique Ghana Card number."},
        {"name": "group_name", "required": True, "help": "Borrower group, cluster, or team name."},
        {"name": "preferred_language", "required": False, "help": "Defaults to the first configured company language when omitted."},
        {"name": "status", "required": False, "help": "Optional borrower status from legacy systems: ACTIVE, CLOSED, or BAD."},
    ]


def infer_borrower_import_format(filename: str) -> str | None:
    lowered = (filename or "").lower()
    if lowered.endswith(".csv"):
        return "csv"
    if lowered.endswith(".xlsx"):
        return "excel"
    if lowered.endswith(".json"):
        return "json"
    if lowered.endswith(".xml"):
        return "xml"
    return None


def parse_borrower_import(uploaded_file, file_type: str) -> list[dict]:
    if file_type == "csv":
        return _parse_csv(uploaded_file)
    if file_type == "excel":
        return _parse_excel(uploaded_file)
    if file_type == "json":
        return _parse_json(uploaded_file)
    if file_type == "xml":
        return _parse_xml(uploaded_file)
    raise ValueError(f"Unsupported borrower import type: {file_type}")


def import_borrower_rows(rows: list[dict], *, import_mode: str = "create_only") -> dict:
    language_choices = [value for value, _label in Borrower.get_language_choices()]
    default_language = language_choices[0] if language_choices else "Twi"

    created = 0
    updated = 0
    errors: list[BorrowerImportError] = []
    touched = []

    for index, raw_row in enumerate(rows, start=2):
        row = _normalize_row(raw_row)
        try:
            borrower, action = _save_borrower_row(
                row,
                import_mode=import_mode,
                default_language=default_language,
                allowed_languages=language_choices,
            )
        except Exception as exc:
            errors.append(BorrowerImportError(index, str(exc), row))
            continue

        touched.append(
            {
                "borrower_id": borrower.borrower_id,
                "full_name": borrower.full_name,
                "action": action,
            }
        )
        if action == "created":
            created += 1
        else:
            updated += 1

    return {
        "total_rows": len(rows),
        "created": created,
        "updated": updated,
        "errors": errors,
        "touched": touched,
    }


def build_borrower_import_sample(file_type: str):
    sample_rows = [
        {
            "borrower_id": "BRW-1001",
            "full_name": "Ama Mensah",
            "national_id": "GHA-AMA-1001",
            "phone": "0241000001",
            "email": "ama.mensah@example.com",
            "location": "Accra Central",
            "ghana_card_number": "GHA-123456789-1",
            "group_name": "Makola Traders",
            "preferred_language": "Twi",
        },
        {
            "borrower_id": "BRW-1002",
            "full_name": "Kojo Asare",
            "national_id": "GHA-KOJO-1002",
            "phone": "0241000002",
            "email": "kojo.asare@example.com",
            "location": "Tema Community 1",
            "ghana_card_number": "GHA-123456789-2",
            "group_name": "Tema Retail Group",
            "preferred_language": "Ga",
        },
    ]

    if file_type == "csv":
        stream = io.StringIO()
        writer = csv.DictWriter(stream, fieldnames=BORROWER_SAMPLE_COLUMNS)
        writer.writeheader()
        writer.writerows(sample_rows)
        return ("borrower_import_sample.csv", "text/csv", stream.getvalue().encode("utf-8"))

    if file_type == "excel":
        book = Workbook()
        sheet = book.active
        sheet.title = "Borrowers"
        sheet.append(BORROWER_SAMPLE_COLUMNS)
        for row in sample_rows:
            sheet.append([row.get(column, "") for column in BORROWER_SAMPLE_COLUMNS])
        output = io.BytesIO()
        book.save(output)
        return (
            "borrower_import_sample.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            output.getvalue(),
        )

    if file_type == "json":
        payload = json.dumps(sample_rows, indent=2)
        return ("borrower_import_sample.json", "application/json", payload.encode("utf-8"))

    if file_type == "xml":
        root = ET.Element("borrowers")
        for row in sample_rows:
            node = ET.SubElement(root, "borrower")
            for column in BORROWER_SAMPLE_COLUMNS:
                child = ET.SubElement(node, column)
                child.text = str(row.get(column, ""))
        payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        return ("borrower_import_sample.xml", "application/xml", payload)

    raise ValueError(f"Unsupported sample type: {file_type}")


def _parse_csv(uploaded_file) -> list[dict]:
    uploaded_file.seek(0)
    text = uploaded_file.read().decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [dict(row) for row in reader if any(str(value or "").strip() for value in row.values())]


def _parse_excel(uploaded_file) -> list[dict]:
    uploaded_file.seek(0)
    workbook = load_workbook(uploaded_file, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [_normalize_header(cell) for cell in rows[0]]
    data = []
    for values in rows[1:]:
        if not any(value not in (None, "") for value in values):
            continue
        item = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            item[header] = values[index] if index < len(values) else ""
        data.append(item)
    return data


def _parse_json(uploaded_file) -> list[dict]:
    uploaded_file.seek(0)
    payload = json.loads(uploaded_file.read().decode("utf-8"))
    if isinstance(payload, dict):
        payload = payload.get("borrowers", [])
    if not isinstance(payload, list):
        raise ValueError("JSON import must be a list of borrower objects or a {'borrowers': [...]} payload.")
    return [dict(item) for item in payload if isinstance(item, dict)]


def _parse_xml(uploaded_file) -> list[dict]:
    uploaded_file.seek(0)
    root = ET.fromstring(uploaded_file.read())
    if root.tag.lower() == "borrower":
        nodes = [root]
    else:
        nodes = [child for child in root if child.tag.lower() == "borrower"]
    rows = []
    for node in nodes:
        row = {}
        for child in node:
            row[_normalize_header(child.tag)] = child.text or ""
        if any(str(value or "").strip() for value in row.values()):
            rows.append(row)
    return rows


def _normalize_header(value) -> str:
    return (
        str(value or "")
        .strip()
        .lower()
        .replace(" ", "_")
        .replace("-", "_")
    )


def _normalize_row(row: dict) -> dict:
    normalized = {}
    for key, value in row.items():
        header = _normalize_header(key)
        if not header:
            continue
        normalized[header] = _stringify(value)
    return normalized


def _stringify(value) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return str(value)
    return str(value).strip()


def _save_borrower_row(row: dict, *, import_mode: str, default_language: str, allowed_languages: list[str]):
    borrower_id = row.get("borrower_id", "")
    if not borrower_id:
        raise ValueError("Missing required field: borrower_id.")

    full_name = row.get("full_name", "")
    national_id = row.get("national_id", "")
    phone = row.get("phone", "")
    location = row.get("location", "")
    ghana_card_number = row.get("ghana_card_number", "")
    group_name = row.get("group_name", "")

    required_text = {
        "full_name": full_name,
        "national_id": national_id,
        "phone": phone,
        "location": location,
        "ghana_card_number": ghana_card_number,
        "group_name": group_name,
    }
    for field_name, value in required_text.items():
        if not value:
            raise ValueError(f"Missing required field: {field_name}.")

    existing = Borrower.objects.filter(borrower_id=borrower_id).first()
    if existing is not None and import_mode == "create_only":
        raise ValueError(f"Borrower ID '{borrower_id}' already exists.")
    if existing is None:
        limit_check = can_create_borrower()
        if not limit_check.allowed:
            raise ValueError(limit_check.reason)

    target = existing or Borrower()

    preferred_language = row.get("preferred_language") or default_language
    if preferred_language not in allowed_languages:
        raise ValueError(
            f"Unsupported preferred_language '{preferred_language}'. Allowed values: {', '.join(allowed_languages)}."
        )

    email = row.get("email", "").lower()
    _validate_import_email(email, exclude_user_id=target.user_id)

    _assert_unique_borrower_fields(target, national_id=national_id, ghana_card_number=ghana_card_number)

    target.borrower_id = borrower_id
    target.full_name = full_name
    target.national_id = national_id
    target.phone = phone
    target.email = email
    target.location = location
    target.ghana_card_number = ghana_card_number
    target.group_name = group_name
    target.preferred_language = preferred_language

    status = (row.get("status") or "").upper()
    if status in dict(Borrower.STATUS_CHOICES):
        target.status = status

    with transaction.atomic():
        target.save()

    return target, "updated" if existing else "created"


def _validate_import_email(email: str, *, exclude_user_id=None):
    if not email:
        return
    users = UserModel._default_manager.filter(email__iexact=email)
    if exclude_user_id:
        users = users.exclude(pk=exclude_user_id)
    if users.exists():
        raise ValueError(f"Email '{email}' is already linked to another login account.")


def _assert_unique_borrower_fields(target: Borrower, *, national_id: str, ghana_card_number: str):
    national_match = Borrower.objects.filter(national_id=national_id)
    card_match = Borrower.objects.filter(ghana_card_number=ghana_card_number)
    if target.pk:
        national_match = national_match.exclude(pk=target.pk)
        card_match = card_match.exclude(pk=target.pk)

    if national_match.exists():
        raise ValueError(f"National ID '{national_id}' already belongs to another borrower.")
    if card_match.exists():
        raise ValueError(f"Ghana Card number '{ghana_card_number}' already belongs to another borrower.")
