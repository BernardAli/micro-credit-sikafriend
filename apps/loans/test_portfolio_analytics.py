from datetime import timedelta
from decimal import Decimal
from io import BytesIO
import csv
import zipfile
from unittest.mock import patch

from django.contrib.auth.models import User
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import Agent
from .models import Borrower, Loan, Payment, Portfolio, PortfolioAssignment, RepaymentCollection, SecurityDeposit, SecurityDepositTransaction, SecurityDepositTransfer
from .portfolio_analytics import portfolio_snapshot, visible_portfolios, dashboard_portfolios
from .portfolios import assign_customer, set_portfolio_agents
from .portfolio_exports import summary_row


class PortfolioAnalyticsTests(TestCase):
    def setUp(self):
        self.today = timezone.localdate()
        self.start = self.today - timedelta(days=90)
        self.portfolio = Portfolio.objects.get(number="07")
        self.other = Portfolio.objects.get(number="01")
        self.admin = User.objects.create_user("analytics-manager", is_staff=True)
        self.agent_user = User.objects.create_user("analytics-officer")
        self.agent = Agent.objects.create(user=self.agent_user, agent_id="AN-1", name="Portfolio Officer", phone="0241111111", location="Accra")
        set_portfolio_agents(self.portfolio, [self.agent], actor=self.admin)
        self.customer = self.new_customer("Ama", self.portfolio)
        self.other_customer = self.new_customer("PRIVATE-OTHER", self.other)
        self.loan = self.new_loan(self.customer)
        notify = patch("apps.loans.signals.notify_user")
        notify.start()
        self.addCleanup(notify.stop)
        self.payment = Payment.objects.create(loan=self.loan, amount="60", payment_date=self.today, submitted_by=self.agent_user)
        Payment.objects.create(loan=self.loan, amount="500", payment_date=self.today, is_pending=True)
        Payment.objects.create(loan=self.loan, amount="30", payment_date=self.today, is_void=True)
        self.client.force_login(self.admin)

    def new_customer(self, name, portfolio):
        return Borrower.objects.create(portfolio=portfolio, full_name=name, national_id=name, ghana_card_number=name,
                                       phone="0242222222", location="Accra", group_name="Traders", preferred_language="Twi")

    def new_loan(self, customer):
        loan = Loan(borrower=customer, agent_name=self.agent, amount="1000", duration_weeks=10, weekly_payment="120", start_date=self.today - timedelta(days=14))
        loan._skip_new_loan_notifications = True
        loan.save()
        return loan

    def snapshot(self):
        return portfolio_snapshot(Portfolio.objects.filter(pk=self.portfolio.pk), self.start, self.today)

    def test_financial_totals_exclude_pending_and_void_payments(self):
        metrics = self.snapshot()["totals"]
        self.assertEqual(metrics["period_collected"], Decimal("60.00"))
        self.assertEqual(metrics["outstanding"], Decimal("1140.00"))
        self.assertEqual(metrics["period_payments"], 1)
        self.assertEqual(metrics["pending_amount"], Decimal("500.00"))
        self.assertEqual(metrics["void_payments"], 1)

    def test_partial_payment_does_not_clear_an_installment(self):
        row = self.snapshot()["books"][self.portfolio.pk]["loans"][0]
        self.assertEqual(row["arrears"], Decimal("180.00"))
        self.assertEqual(row["days_overdue"], 7)
        self.assertEqual(row["missed_installments"], 2)

    def test_par30_uses_principal_and_excludes_undisbursed_loans(self):
        self.loan.start_date = self.today - timedelta(days=50)
        self.loan.save()
        pending = self.new_loan(self.customer)
        SecurityDeposit.objects.create(loan=pending, borrower=self.customer, required_amount="100")
        metrics = self.snapshot()["totals"]
        self.assertEqual(metrics["par30_principal"], Decimal("950.00"))
        self.assertEqual(metrics["risk_principal"], Decimal("950.00"))
        self.assertEqual(metrics["par30_rate"], Decimal("100.00"))
        self.assertEqual(metrics["awaiting_disbursement"], 1)

    def test_security_deposits_are_separate_from_repayments(self):
        loan = self.new_loan(self.customer)
        deposit = SecurityDeposit.objects.create(loan=loan, borrower=self.customer, required_amount="100")
        SecurityDepositTransaction.objects.create(deposit=deposit, action="COLLECTION", amount="80", reference="DEP-80")
        data = self.snapshot()
        self.assertEqual(data["totals"]["deposit_held"], Decimal("80.00"))
        self.assertEqual(data["totals"]["deposit_due"], Decimal("20.00"))
        self.assertEqual(data["totals"]["period_collected"], Decimal("60.00"))
        self.assertEqual(next(row for row in data["books"][self.portfolio.pk]["loans"] if row["pk"] == loan.pk)["arrears"], Decimal("0.00"))

    def test_activity_period_does_not_change_current_balances(self):
        data = portfolio_snapshot(Portfolio.objects.filter(pk=self.portfolio.pk), self.today - timedelta(days=100), self.today - timedelta(days=95))
        self.assertEqual(data["totals"]["period_collected"], Decimal("0.00"))
        self.assertEqual(data["totals"]["outstanding"], Decimal("1140.00"))

    def test_all_detail_sections_and_dashboard_navigation(self):
        response = self.client.get(reverse("loans:dashboard"))
        self.assertContains(response, "Explore all portfolios")
        self.assertContains(response, reverse("loans:portfolio_detail", args=[self.portfolio.pk]))
        for tab in ("overview", "customers", "loans", "payments", "risk", "deposits", "collections", "officers", "ids"):
            response = self.client.get(reverse("loans:portfolio_detail", args=[self.portfolio.pk]), {"tab": tab})
            self.assertEqual(response.status_code, 200, tab)

    def test_agent_pages_and_exports_are_scoped_to_assigned_portfolios(self):
        self.client.force_login(self.agent_user)
        response = self.client.get(reverse("loans:portfolio_list"), {"export": "json"})
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "PRIVATE-OTHER")
        for fmt in (None, "json", "excel", "csv", "pdf"):
            response = self.client.get(reverse("loans:portfolio_detail", args=[self.other.pk]), {"export": fmt} if fmt else {})
            self.assertEqual(response.status_code, 404)
        set_portfolio_agents(self.portfolio, [])
        self.assertEqual(self.client.get(reverse("loans:portfolio_detail", args=[self.portfolio.pk])).status_code, 404)

    def test_overview_filters_apply_to_exports(self):
        response = self.client.get(reverse("loans:portfolio_list"), {"q": "Wisdom", "export": "json"})
        rows = response.json()["tables"][0]["rows"]
        self.assertEqual([row["number"] for row in rows], ["07"])

    def test_export_includes_rows_beyond_screen_pagination(self):
        for number in range(30):
            self.new_customer(f"Customer-{number}", self.portfolio)
        url = reverse("loans:portfolio_detail", args=[self.portfolio.pk])
        response = self.client.get(url, {"tab": "customers"})
        self.assertEqual(len(response.context["page_obj"]), 25)
        response = self.client.get(url, {"export": "csv", "section": "customers", "page": "2"})
        rows = list(csv.DictReader(response.content.decode("utf-8-sig").splitlines()))
        self.assertEqual(len(rows), 31)
        self.assertTrue(any(row["Customer"] == "Customer-29" for row in rows))

    def test_workbook_json_and_csv_bundle_contain_complete_registers(self):
        url = reverse("loans:portfolio_detail", args=[self.portfolio.pk])
        response = self.client.get(url, {"export": "excel"})
        workbook = load_workbook(BytesIO(response.content))
        self.assertIn("Customer Register", workbook.sheetnames)
        self.assertIn("Repayment Aging", workbook.sheetnames)
        self.assertEqual(workbook["Customer Register"].max_row, 2)
        response = self.client.get(url, {"export": "csv"})
        with zipfile.ZipFile(BytesIO(response.content)) as archive:
            self.assertIn("customers.csv", archive.namelist())
            self.assertIn("report-notes.txt", archive.namelist())
        response = self.client.get(url, {"export": "json", "section": "loans"})
        self.assertEqual(response.json()["tables"][0]["rows"][0]["customer_id"], self.customer.borrower_id)

    def test_pdf_report_is_a_real_download(self):
        response = self.client.get(reverse("loans:portfolio_detail", args=[self.portfolio.pk]), {"export": "pdf", "section": "summary"})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"%PDF"))
        self.assertIn("attachment", response["Content-Disposition"])

    def test_search_export_and_csv_formula_safety(self):
        self.customer.full_name = "=2+2"
        self.customer.save()
        response = self.client.get(reverse("loans:portfolio_detail", args=[self.portfolio.pk]), {"export": "csv", "section": "customers", "q": self.customer.borrower_id})
        self.assertIn("'=2+2", response.content.decode())

    def test_multiple_officers_do_not_duplicate_portfolio_totals(self):
        second = Agent.objects.create(agent_id="AN-2", name="Second Officer", phone="0240000000", location="Accra")
        set_portfolio_agents(self.portfolio, [self.agent, second])
        self.assertEqual(self.snapshot()["totals"]["outstanding"], Decimal("1140.00"))
        cards = dashboard_portfolios(self.admin)
        self.assertEqual(next(row for row in cards["rows"] if row.pk == self.portfolio.pk).loan_count, 1)

    def test_invalid_dates_or_export_sections_fail_cleanly(self):
        url = reverse("loans:portfolio_detail", args=[self.portfolio.pk])
        for params in ({"date_from": "invalid"}, {"date_from": "2026-10-11", "date_to": "2026-10-01"}, {"export": "invalid"}, {"export": "csv", "section": "invalid"}):
            self.assertEqual(self.client.get(url, params).status_code, 400)

    def test_transfers_move_exposure_and_preserve_id(self):
        customer_id = self.customer.borrower_id
        assign_customer(self.customer, self.other, actor=self.admin)
        data = self.snapshot()
        self.assertEqual(data["totals"]["outstanding"], Decimal("0.00"))
        self.assertEqual(data["books"][self.portfolio.pk]["issuance"][0]["customer_id"], customer_id)
        self.assertEqual(data["books"][self.portfolio.pk]["issuance"][0]["customer"], "Former customer")
        data = portfolio_snapshot(Portfolio.objects.filter(pk=self.other.pk), self.start, self.today)
        self.assertEqual(data["totals"]["outstanding"], Decimal("1140.00"))

    def test_customer_population_includes_customers_without_loans(self):
        self.new_customer("No Loan", self.portfolio)
        self.assertEqual(self.snapshot()["totals"]["customers_without_loans"], 1)

    def test_snapshot_queries_do_not_grow_per_customer_or_loan(self):
        for number in range(8):
            self.new_loan(self.new_customer(f"Scale-{number}", self.portfolio))
        with CaptureQueriesContext(connection) as queries:
            self.snapshot()
        self.assertLessEqual(len(queries), 18)

    def test_payment_phone_search_matches_customer(self):
        response = self.client.get(reverse("loans:portfolio_detail", args=[self.portfolio.pk]), {"export": "json", "section": "payments", "q": self.customer.phone})
        self.assertEqual(len(response.json()["tables"][0]["rows"]), 3)

    def test_collection_attempts_do_not_double_count_repayments(self):
        RepaymentCollection.objects.create(loan=self.loan, borrower=self.customer, payment=self.payment, amount="60", client_reference="RECONCILED-1", status="RECONCILED")
        RepaymentCollection.objects.create(loan=self.loan, borrower=self.customer, amount="90", client_reference="REVIEW-1", status="MANUAL_REVIEW", failure_reason="Reference needs verification")
        metrics = self.snapshot()["totals"]
        self.assertEqual(metrics["period_collected"], Decimal("60.00"))
        self.assertEqual(metrics["pipeline_amount"], Decimal("90.00"))
        self.assertEqual(metrics["review_collections"], 1)

    def test_summary_matches_full_totals_with_credit_deposits_and_pipeline(self):
        excess = Payment(loan=self.loan, amount=Decimal('1500'))
        excess._legacy_import = True
        excess.save()
        other_loan = self.new_loan(self.customer)
        source = SecurityDeposit.objects.create(loan=self.loan, borrower=self.customer, required_amount=Decimal('100'))
        target = SecurityDeposit.objects.create(loan=other_loan, borrower=self.customer, required_amount=Decimal('100'))
        for action, amount in [('COLLECTION', '100'), ('COLLECTION', '100'), ('REFUND', '20'), ('DEFAULT_APPLICATION', '10')]:
            SecurityDepositTransaction.objects.create(deposit=source, action=action, amount=Decimal(amount))
        SecurityDepositTransfer.objects.create(source=source, target=target, amount=Decimal('100'), reference='SUMMARY-TRANSFER')
        old_review = RepaymentCollection.objects.create(loan=self.loan, borrower=self.customer, amount='90', client_reference='SUMMARY-OLD-REVIEW', status='MANUAL_REVIEW')
        old_failed = RepaymentCollection.objects.create(loan=self.loan, borrower=self.customer, amount='20', client_reference='SUMMARY-OLD-FAILED', status='FAILED')
        RepaymentCollection.objects.filter(pk__in=[old_review.pk, old_failed.pk]).update(initiated_at=timezone.now()-timedelta(days=120))
        RepaymentCollection.objects.create(loan=self.loan, borrower=self.customer, payment=self.payment, amount='60', client_reference='SUMMARY-RECONCILED', status='RECONCILED')
        PortfolioAssignment.objects.create(portfolio=self.portfolio, agent=self.agent, ended_at=timezone.now(), assigned_by=self.admin)
        full = self.snapshot()
        summary = portfolio_snapshot(Portfolio.objects.filter(pk=self.portfolio.pk), self.start, self.today, summary_only=True)
        self.assertEqual(summary['totals'], full['totals'])
        self.assertEqual(summary_row(summary['books'][self.portfolio.pk]), summary_row(full['books'][self.portfolio.pk]))
        self.assertEqual(full['totals']['customer_credit'], Decimal('360'))
        self.assertEqual(full['totals']['deposit_held'], Decimal('170'))
        self.assertEqual(full['totals']['pipeline_amount'], Decimal('90'))
        self.assertEqual(full['totals']['collection_attempts'], 2)
        self.assertEqual(full['totals']['current_pending_amount'], Decimal('500'))

    def test_assignment_csv_query_count_stays_constant_for_more_customers(self):
        url = reverse('loans:portfolio_export')
        with CaptureQueriesContext(connection) as initial:
            self.client.get(url)
        for number in range(15):
            self.new_customer(f'CSV-Scale-{number}', self.portfolio)
        with CaptureQueriesContext(connection) as expanded:
            response = self.client.get(url)
        self.assertLessEqual(len(expanded), len(initial))
        self.assertLessEqual(len(expanded), 3)
        rows = list(csv.DictReader(response.content.decode().splitlines()))
        self.assertEqual(len(rows), Borrower.objects.count())
        own_row = next(row for row in rows if row['borrower_pk']==str(self.customer.pk))
        self.assertEqual(own_row['customer_id'], self.customer.borrower_id)
        self.assertEqual(own_row['portfolio_number'], self.portfolio.number)

    def test_rendered_portfolio_pages_queries_do_not_grow_per_loan(self):
        pages = [
            (reverse('loans:portfolio_list'), {}),
            (reverse('loans:portfolio_detail', args=[self.portfolio.pk]), {'tab':'payments'}),
            (reverse('loans:portfolio_detail', args=[self.portfolio.pk]), {'tab':'deposits'}),
            (reverse('loans:portfolio_detail', args=[self.portfolio.pk]), {'tab':'officers'}),
            (reverse('loans:portfolio_customers'), {'show':'all'}),
            (reverse('loans:portfolio_update', args=[self.portfolio.pk]), {}),
        ]
        counts = []
        for url, params in pages:
            self.client.get(url, params)
            with CaptureQueriesContext(connection) as queries:
                self.client.get(url, params)
            counts.append(len(queries))
        for number in range(10):
            loan = self.new_loan(self.new_customer(f'Page-Scale-{number}', self.portfolio))
            deposit = SecurityDeposit.objects.create(loan=loan, borrower=loan.borrower, required_amount=Decimal('50'))
            SecurityDepositTransaction.objects.create(deposit=deposit, action='COLLECTION', amount=Decimal('50'))
            Payment.objects.create(loan=loan, amount=Decimal('120'))
        for (url, params), before in zip(pages, counts):
            with self.subTest(page=url, params=params):
                with CaptureQueriesContext(connection) as queries:
                    response = self.client.get(url, params)
                self.assertEqual(response.status_code, 200)
                self.assertLessEqual(len(queries), before)
