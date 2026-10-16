"""Bound import work before decoding archives or creating application records."""
from functools import wraps
from zipfile import BadZipFile, ZipFile
from django.conf import settings
from defusedxml import ElementTree
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook


def limited_rows(rows):
    result = []
    for row in rows:
        if len(result) >= settings.IMPORT_MAX_ROWS:
            raise ValueError(f"Import at most {settings.IMPORT_MAX_ROWS:,} rows per file.")
        result.append(row)
    return result


def bounded_import(parser):
    @wraps(parser)
    def checked(uploaded_file, file_type):
        uploaded_file.seek(0, 2)
        size = uploaded_file.tell()
        uploaded_file.seek(0)
        if size > settings.IMPORT_MAX_BYTES:
            raise ValueError("Import file exceeds the permitted size. Split it into smaller files.")
        if file_type == 'excel':
            try:
                with ZipFile(uploaded_file) as archive:
                    members = archive.infolist()
                    if len(members) > 2048 or sum(member.file_size for member in members) > settings.IMPORT_MAX_EXPANDED_BYTES:
                        raise ValueError("Spreadsheet expands beyond the permitted size.")
            except BadZipFile:
                raise ValueError("Upload a valid .xlsx spreadsheet.") from None
            uploaded_file.seek(0)
        elif file_type == 'json':
            # Reject excessive containers before json.loads can allocate them.
            text = uploaded_file.read().decode('utf-8-sig')
            quoted = escaped = False
            depth = containers = 0
            for character in text:
                if quoted:
                    if escaped:
                        escaped = False
                    elif character == '\\':
                        escaped = True
                    elif character == '"':
                        quoted = False
                elif character == '"':
                    quoted = True
                elif character in '[{':
                    depth += 1; containers += 1
                    if depth > 20 or containers > settings.IMPORT_MAX_ROWS * 3:
                        raise ValueError("JSON import is too large or deeply nested.")
                elif character in ']}':
                    depth -= 1
            uploaded_file.seek(0)
        return limited_rows(parser(uploaded_file, file_type))
    return checked


def safe_xml(data):
    try:
        return ElementTree.fromstring(data, forbid_dtd=True, forbid_entities=True, forbid_external=True)
    except DefusedXmlException:
        raise ValueError("XML imports must not contain DTDs, entities or external references.") from None


def excel_rows(uploaded_file):
    workbook = load_workbook(uploaded_file, data_only=True, read_only=True, keep_links=False)
    try:
        sheet = workbook.active
        if (sheet.max_column or 0) > settings.IMPORT_MAX_COLUMNS:
            raise ValueError("Spreadsheet contains too many columns.")
        return limited_rows_with_header(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()


def limited_rows_with_header(rows):
    result = []
    for row in rows:
        if len(row) > settings.IMPORT_MAX_COLUMNS or len(result) > settings.IMPORT_MAX_ROWS:
            raise ValueError("Spreadsheet exceeds the row or column limit.")
        result.append(row)
    return result
